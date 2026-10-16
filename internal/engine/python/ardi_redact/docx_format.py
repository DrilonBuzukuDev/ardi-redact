"""DOCX adapter preserves package contents and edits supported Word text nodes."""

import re
from collections.abc import Iterator
from dataclasses import dataclass
from io import BytesIO
from typing import BinaryIO
from zipfile import ZipFile, ZipInfo

from docx import Document
from lxml import etree
from lxml.etree import _Element as XmlElement  # pyright: ignore[reportPrivateUsage]

from ardi_redact.redaction import TextUnit, apply_edits

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


@dataclass
class XmlParagraph:
  nodes: list[tuple[XmlElement, int, str]]


class DocxFile:
  def __init__(self, data: bytes) -> None:
    # python-docx checks the Word package. Modify text nodes rather than
    # assigning Run.text, which would discard drawings and other run content.
    Document(BytesIO(data))
    self.units: list[TextUnit] = []
    self.warnings: list[str] = []
    self.paragraphs: list[XmlParagraph] = []
    self.entries: list[tuple[ZipInfo, bytes]] = []
    self.roots: dict[str, XmlElement] = {}
    with ZipFile(BytesIO(data)) as archive:
      self.comment = archive.comment
      for info in archive.infolist():
        content = archive.read(info)
        self.entries.append((info, content))
        name = info.filename
        if name == "word/document.xml" or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name):
          root = etree.fromstring(content, etree.XMLParser(resolve_entities=False))
          self.roots[name] = root
          self._paragraphs(root)
        elif name in {"word/comments.xml", "word/footnotes.xml", "word/endnotes.xml"}:
          self.warnings.append("DOCX comments, footnotes, and endnotes are not scanned.")
        elif name.startswith(("word/media/", "word/embeddings/")):
          self.warnings.append("DOCX images and embedded objects are not scanned.")
        elif name.endswith(".rels") and b'TargetMode="External"' in content:
          self.warnings.append("DOCX external hyperlink targets are not scanned.")
    self.warnings.append("DOCX document properties and other package metadata are not scanned.")
    self.warnings = sorted(set(self.warnings))

  def _paragraphs(self, root: XmlElement) -> None:
    if next(root.iter(W + "txbxContent"), None) is not None:
      self.warnings.append("DOCX text boxes are not scanned.")
    if next(root.iter(W + "instrText"), None) is not None:
      self.warnings.append("DOCX field instructions are not scanned; displayed text is scanned.")
    if next(root.iter(W + "del"), None) is not None:
      self.warnings.append("DOCX deleted tracked changes are not scanned.")
    if next(root.iter(W + "altChunk"), None) is not None:
      self.warnings.append("DOCX imported alternative-format content is not scanned.")
    for paragraph in root.iter(W + "p"):
      if any(parent.tag == W + "txbxContent" for parent in paragraph.iterancestors()):
        continue
      source = ""
      nodes: list[tuple[XmlElement, int, str]] = []
      protected: set[int] = set()
      for node in self._text_nodes(paragraph):
        if node.tag == W + "t":
          value = node.text or ""
          nodes.append((node, len(source), value))
          source += value
        else:
          protected.add(len(source))
          source += "\n" if node.tag in {W + "br", W + "cr"} else "\t"
      if nodes:
        # Tabs and line breaks are represented by whitespace for NER.
        unit = TextUnit.project(source)
        unit.offsets = [None if index in protected else index for index in unit.offsets]
        self.units.append(unit)
        self.paragraphs.append(XmlParagraph(nodes))

  def _text_nodes(self, paragraph: XmlElement) -> Iterator[XmlElement]:
    for node in paragraph.iter():
      if node.tag not in {W + "t", W + "tab", W + "br", W + "cr"}:
        continue
      # Nested text boxes have their own paragraphs and are unsupported.
      if any(parent.tag in {W + "txbxContent", W + "del"} for parent in node.iterancestors()):
        continue
      yield node

  def write(self, edits: list[dict[int, str]], stream: BinaryIO) -> None:
    for paragraph, changes in zip(self.paragraphs, edits, strict=True):
      for node, start, original in paragraph.nodes:
        updated = apply_edits(original, changes, start)
        node.text = updated
        if updated.startswith(" ") or updated.endswith(" "):
          node.set(XML_SPACE, "preserve")
    with ZipFile(stream, "w") as archive:
      archive.comment = self.comment
      for info, content in self.entries:
        if info.filename in self.roots:
          content = etree.tostring(
            self.roots[info.filename],
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
          )
        archive.writestr(info, content)
